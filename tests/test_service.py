from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from mdforge.converters.pdf import _extract_page
from mdforge.converters.pptx import PptxConverter
from mdforge.converters.table import table_to_markdown
from mdforge.converters.xlsx import XlsxConverter
from mdforge.models import ConversionOptions
from mdforge.service import ConversionService


def _make_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Vendas"
    ws1.append(["Produto", "Total"])
    ws1.append(["Caneta", 10])
    ws2 = wb.create_sheet("Resumo")
    ws2.append(["Item", "Valor"])
    ws2.append(["X", 1])
    wb.save(path)
    return path


def _word(text: str, x0: float, x1: float, top: float) -> dict:
    return {"text": text, "x0": x0, "x1": x1, "top": top, "bottom": top + 10}


class _FakePdfPage:
    def __init__(self, words: list[dict]):
        self.words = words

    def extract_words(self, **kwargs):
        return self.words

    def extract_text(self, **kwargs):
        return "fallback"


def test_txt_to_md(tmp_path: Path):
    source = tmp_path / "hello.txt"
    source.write_text("Olá mundo", encoding="utf-8")
    out = tmp_path / "out"
    result = ConversionService().convert_file(source, out, ConversionOptions(False, False))
    assert result.success
    assert (out / "hello.md").read_text(encoding="utf-8") == "Olá mundo\n"


def test_unsupported_format(tmp_path: Path):
    source = tmp_path / "sample.xyz"
    source.write_text("x", encoding="utf-8")
    result = ConversionService().convert_file(source, tmp_path / "out")
    assert not result.success


def test_custom_output_name(tmp_path: Path):
    source = tmp_path / "hello.txt"
    source.write_text("oi", encoding="utf-8")
    out = tmp_path / "out"
    options = ConversionOptions(include_source_header=False, output_name="renomeado.md")
    result = ConversionService().convert_file(source, out, options)
    assert result.success
    assert result.destination == (out / "renomeado.md").resolve()
    assert (out / "renomeado.md").exists()


def test_output_name_ignored_for_multiple_files(tmp_path: Path):
    a = tmp_path / "a.txt"
    b = tmp_path / "b.txt"
    a.write_text("a", encoding="utf-8")
    b.write_text("b", encoding="utf-8")
    out = tmp_path / "out"
    options = ConversionOptions(include_source_header=False, output_name="unico.md")
    results = ConversionService().convert_many([a, b], out, options)
    assert all(r.success for r in results)
    assert (out / "a.md").exists()
    assert (out / "b.md").exists()


def test_xlsx_all_sheets(tmp_path: Path):
    source = _make_xlsx(tmp_path / "planilha.xlsx")
    out = tmp_path / "out"
    result = ConversionService().convert_file(
        source, out, ConversionOptions(include_source_header=False)
    )
    assert result.success
    content = (out / "planilha.md").read_text(encoding="utf-8")
    assert "## Vendas" in content
    assert "## Resumo" in content
    assert "| Produto | Total |" in content


def test_xlsx_single_sheet(tmp_path: Path):
    source = _make_xlsx(tmp_path / "planilha.xlsx")
    out = tmp_path / "out"
    options = ConversionOptions(include_source_header=False, sheet_name="Resumo")
    result = ConversionService().convert_file(source, out, options)
    assert result.success
    content = (out / "planilha.md").read_text(encoding="utf-8")
    assert "## Resumo" in content
    assert "## Vendas" not in content


def test_xlsx_unknown_sheet_fails(tmp_path: Path):
    source = _make_xlsx(tmp_path / "planilha.xlsx")
    out = tmp_path / "out"
    options = ConversionOptions(include_source_header=False, sheet_name="Inexistente")
    result = ConversionService().convert_file(source, out, options)
    assert not result.success
    assert "Inexistente" in result.message


def test_xlsx_sheet_names(tmp_path: Path):
    source = _make_xlsx(tmp_path / "planilha.xlsx")
    assert XlsxConverter().sheet_names(source) == ["Vendas", "Resumo"]


def test_shared_table_formatter_handles_caption_and_pipes():
    rows = [
        ["Cálculo do custo de capital próprio (Ke)", "", ""],
        ["", "Parâmetros", "Notas"],
        ["Taxa livre de risco | Rf EUA", "1.54%", "[a]"],
        ["Ke", "11.04%", "-"],
    ]
    markdown = table_to_markdown(rows, extract_caption=True)
    assert markdown.startswith("**Cálculo do custo de capital próprio (Ke)**")
    assert "|  | Parâmetros | Notas |" in markdown
    assert "Taxa livre de risco \\| Rf EUA" in markdown


def test_pptx_table_is_preserved_as_markdown(tmp_path: Path):
    source = tmp_path / "table.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(
        4, 3, Inches(1), Inches(1), Inches(8), Inches(3)
    ).table

    title = table.cell(0, 0)
    title.merge(table.cell(0, 2))
    title.text = "Cálculo do custo de capital próprio (Ke)"
    table.cell(1, 1).text = "Parâmetros"
    table.cell(1, 2).text = "Notas"
    table.cell(2, 0).text = "Taxa livre de risco (Rf EUA)"
    table.cell(2, 1).text = "1.54%"
    table.cell(2, 2).text = "[a]"
    table.cell(3, 0).text = "Ke"
    table.cell(3, 1).text = "11.04%"
    table.cell(3, 2).text = "-"
    prs.save(source)

    markdown = PptxConverter().convert(source)
    assert "**Cálculo do custo de capital próprio (Ke)**" in markdown
    assert "|  | Parâmetros | Notas |" in markdown
    assert "| Taxa livre de risco (Rf EUA) | 1.54% | [a] |" in markdown


def test_pdf_layout_inference_recovers_three_column_table():
    words = [_word("Cálculo do custo de capital próprio (Ke)", 50, 260, 20)]
    words += [
        _word("Parâmetros", 370, 430, 50),
        _word("Notas", 500, 530, 50),
    ]
    rows = [
        ("Taxa livre de risco (Rf EUA)", "1.54%", "[a]"),
        ("Prêmio de risco de mercado (PRm)", "5.81%", "[b]"),
        ("Beta desalavancado (βu)", "1.27", "[c]"),
        ("Capital de terceiros / Capital próprio da empresa (D/E)", "-", "-"),
        ("Taxa de IR&CS*", "-", "-"),
        ("Beta realavancado (βl)", "1.27", "[d]"),
        ("Prêmio pelo Risco Brasil (Rp)", "2.10%", "[e]"),
        ("Ke (U$ Real): Rf + βl*PRm + Rp", "11.04%", "-"),
        ("Prêmio Liquidez", "3.00%", "[f]"),
        ("* Ke (R$ Real)", "14.04%", "[=]"),
    ]
    for index, (label, value, note) in enumerate(rows, start=1):
        top = 50 + index * 18
        words.extend(
            [
                _word(label, 50, 300, top),
                _word(value, 370, 420, top),
                _word(note, 500, 530, top),
            ]
        )
    words.append(_word("Notas:", 50, 90, 250))

    markdown = _extract_page(_FakePdfPage(words))
    assert "|  | Parâmetros | Notas |" in markdown
    assert "| Taxa livre de risco (Rf EUA) | 1.54% | [a] |" in markdown
    assert "| Taxa de IR&CS* | - | - |" in markdown
    assert "| * Ke (R$ Real) | 14.04% | [=] |" in markdown
