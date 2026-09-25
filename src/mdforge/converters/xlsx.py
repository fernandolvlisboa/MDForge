from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from .base import Converter
from .table import table_to_markdown

if TYPE_CHECKING:
    from ..models import ConversionOptions


def _sheet_to_md(worksheet) -> str:
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    markdown = table_to_markdown(rows)
    return markdown or "*Planilha vazia.*"


class XlsxConverter(Converter):
    extensions = (".xlsx", ".xlsm")

    def sheet_names(self, source: Path) -> list[str]:
        """Return the worksheet titles available in *source*."""
        workbook = load_workbook(str(source), read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def convert(self, source: Path, options: "ConversionOptions | None" = None) -> str:
        sheet_name = options.sheet_name if options else None
        workbook = load_workbook(str(source), read_only=True, data_only=True)
        try:
            if sheet_name:
                if sheet_name not in workbook.sheetnames:
                    available = ", ".join(workbook.sheetnames)
                    raise ValueError(
                        f"Aba '{sheet_name}' não encontrada. Disponíveis: {available}"
                    )
                sheets = [workbook[sheet_name]]
            else:
                sheets = list(workbook.worksheets)

            out: list[str] = []
            for worksheet in sheets:
                out.append(f"## {worksheet.title}")
                out.append(_sheet_to_md(worksheet))
            return "\n\n".join(out).strip() + "\n"
        finally:
            workbook.close()
