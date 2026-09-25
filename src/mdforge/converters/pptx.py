from pathlib import Path
from typing import TYPE_CHECKING

from pptx import Presentation

from .base import Converter
from .table import table_to_markdown

if TYPE_CHECKING:
    from ..models import ConversionOptions


def _pptx_table_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        values: list[str] = []
        for cell in row.cells:
            values.append("" if getattr(cell, "is_spanned", False) else cell.text)
        rows.append(values)
    return rows


class PptxConverter(Converter):
    extensions = (".pptx",)

    def convert(self, source: Path, options: "ConversionOptions | None" = None) -> str:
        prs = Presentation(str(source))
        out: list[str] = []
        for index, slide in enumerate(prs.slides, start=1):
            out.append(f"## Slide {index}")
            shapes = sorted(slide.shapes, key=lambda shape: (shape.top, shape.left))
            for shape in shapes:
                if getattr(shape, "has_table", False):
                    table_md = table_to_markdown(
                        _pptx_table_rows(shape.table), extract_caption=True
                    )
                    if table_md:
                        out.append(table_md)
                elif getattr(shape, "has_text_frame", False) and shape.text.strip():
                    out.append(shape.text.strip())
            out.append("")
        return "\n\n".join(out).strip() + "\n"
